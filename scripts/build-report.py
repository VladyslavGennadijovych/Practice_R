#!/usr/bin/env python3
"""Build the Ukrainian report from saved R results and genuine RStudio screenshots.

Requires python-docx and Pillow. Run from any directory after analysis.R.
Render the resulting DOCX to PDF with LibreOffice or Microsoft Word.
"""
from pathlib import Path
import csv
from datetime import datetime
from zoneinfo import ZoneInfo

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'reports'
OUT.mkdir(exist_ok=True)

def rows(name):
    with (ROOT / 'results/tables' / (name + '.csv')).open(encoding='utf-8', newline='') as f:
        return list(csv.DictReader(f))

def keyed(name):
    return {r['metric']: r['value'] for r in rows(name)}

metrics = keyed('metrics')
diagnostics = keyed('model_diagnostics')
audit = keyed('data_audit')
desc = rows('descriptive_statistics')
coeff = rows('model_coefficients')
corr = rows('correlations')
scenarios = rows('scenario_predictions')
names = {'ozone_ppb': 'Озон, ppb', 'solar_langley': 'Радіація, langley',
         'wind_ms': 'Вітер, м/с', 'temp_c': 'Температура, °C', '(Intercept)': 'Вільний член'}

def num(x, digits=2):
    return f'{float(x):.{digits}f}'.replace('.', ',').replace('-', '−')

def pv(x):
    x = float(x)
    return f'{x:.2e}' if x < .001 else num(x, 4)

doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = section.bottom_margin = Inches(.65)
section.left_margin = section.right_margin = Inches(.75)
section.footer_distance = Inches(.3)
styles = doc.styles
for style_name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Caption']:
    st = styles[style_name]
    st.font.name = 'Noto Serif'
    st.font.color.rgb = RGBColor(0, 0, 0)
    st.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Noto Serif')
styles['Normal'].font.size = Pt(11)
styles['Normal'].paragraph_format.line_spacing = 1.12
styles['Normal'].paragraph_format.space_after = Pt(6)
styles['Title'].font.size = Pt(22)
styles['Title'].paragraph_format.space_after = Pt(12)
styles['Subtitle'].font.size = Pt(14)
styles['Heading 1'].font.size = Pt(15)
styles['Heading 1'].paragraph_format.space_after = Pt(9)
styles['Heading 1'].paragraph_format.space_before = Pt(0)
styles['Heading 2'].font.size = Pt(12)
styles['Heading 2'].paragraph_format.space_before = Pt(9)
styles['Caption'].font.size = Pt(9)
styles['Caption'].font.italic = True
styles['Caption'].paragraph_format.space_after = Pt(7)
doc.core_properties.title = 'Прикладний статистичний аналіз даних мовою R'
doc.core_properties.subject = 'Звіт з виробничої практики за сімома етапами технічного завдання'
doc.core_properties.author = 'Practice_R'

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = footer.add_run()
field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE')
run._r.addnext(field)
footer.style = 'Caption'

def para(text, bold=False):
    p = doc.add_paragraph()
    p.add_run(text).bold = bold
    return p

def heading(text, new=False):
    p = doc.add_heading(text, 1)
    if new: p.paragraph_format.page_break_before = True

def table(headers, data, widths=None, size=10):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    if widths is None: widths = [7 / len(headers)] * len(headers)
    for col, width in zip(t.columns, widths): col.width = Inches(width)
    for i, header in enumerate(headers): t.rows[0].cells[i].text = str(header)
    for data_row in data:
        for cell, value in zip(t.add_row().cells, data_row): cell.text = str(value)
    for ri, row in enumerate(t.rows):
        trpr = row._tr.get_or_add_trPr()
        no_split = OxmlElement('w:cantSplit'); trpr.append(no_split)
        if ri == 0:
            repeat = OxmlElement('w:tblHeader'); trpr.append(repeat)
        for ci, cell in enumerate(row.cells):
            cell.width = Inches(widths[ci])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr = cell._tc.get_or_add_tcPr()
            border = OxmlElement('w:tcBorders')
            for edge in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
                el = OxmlElement('w:' + edge)
                for k, v in {'val':'single','sz':'4','color':'D9D9D9'}.items(): el.set(qn('w:' + k), v)
                border.append(el)
            tcpr.append(border)
            shade = OxmlElement('w:shd')
            shade.set(qn('w:fill'), 'DFEAF0' if ri == 0 else ('F6F7F8' if ri % 2 == 0 else 'FFFFFF'))
            tcpr.append(shade)
            margins = OxmlElement('w:tcMar')
            for edge in ['top', 'bottom', 'left', 'right']:
                el=OxmlElement('w:' + edge); el.set(qn('w:w'), '75'); el.set(qn('w:type'), 'dxa'); margins.append(el)
            tcpr.append(margins)
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(1)
                p.paragraph_format.line_spacing = 1.0
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT if ci == 0 else WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs:
                    r.font.name = 'Noto Serif'; r.font.size = Pt(size); r.font.bold = ri == 0
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return t

fig_count = 0
def picture(path, caption, width=6.65):
    global fig_count
    fig_count += 1
    p=doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    inline=p.add_run().add_picture(str(path), width=Inches(width))
    inline._inline.docPr.set('descr', caption)
    cp=doc.add_paragraph(f'Рисунок {fig_count}. {caption}', style='Caption')
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER

def figure(name, caption, width=6.65):
    picture(ROOT/'results/figures'/f'{name}.png', caption, width)

def code(text):
    p=doc.add_paragraph()
    p.paragraph_format.line_spacing=1.0
    for i, line in enumerate(text.splitlines()):
        r=p.add_run(('\n' if i else '') + line)
        r.font.name='DejaVu Sans Mono'; r.font.size=Pt(9)

# Page 1: scope, actual result and a concise requirements map.
doc.add_paragraph('Звіт з виробничої практики', style='Title')
doc.add_paragraph('Прикладний статистичний аналіз даних мовою R', style='Subtitle')
para('Тема дослідження — зв’язок концентрації озону з температурою, швидкістю вітру та сонячною радіацією за щоденними спостереженнями в Нью-Йорку у травні–вересні 1973 року.')
para('Розроблено відтворюваний модуль, який читає CSV, перевіряє й очищує дані, формує статистичні таблиці та графіки, будує лінійну модель і перевіряє її на пізніших спостереженнях. Код виконано в R 4.5.3 та RStudio Server 2026.10.0+233.')
para('Основний результат: зі 153 записів отримано 111 повних спостережень. Модель пояснює 61,97% варіації навчальної вибірки, однак тестовий R² становить −2,678. Вона відображає статистичні зв’язки, але потребує доопрацювання для надійного прогнозування.', bold=True)
doc.add_heading('Відповідність технічному завданню', 2)
table(['Етап', 'Виконана робота'], [
    ['1 Підготовка', 'CSV зі 153 записами; аудит пропусків, дублікатів, числових значень і дат; 111 очищених рядків.'],
    ['2 Статистика', 'Середнє, медіана, всі емпіричні моди, дисперсія, стандартне відхилення, мінімум і максимум.'],
    ['3 Візуалізація', 'Гістограми, boxplot, scatterplot та графіки перевірки моделі.'],
    ['4 Кореляція', 'Пірсон, Спірмен, довірчі інтервали, p-значення та корекція BH.'],
    ['5 Регресія', 'Множинна лінійна модель озону; коефіцієнти, інтервали та три сценарії прогнозу.'],
    ['6 Перевірка', 'R², F- і t-статистики, хронологічний тест, залишки, впливові спостереження.'],
    ['7 Висновки', 'Інтерпретація залежностей та обмежень прогнозування.']
], [1.35,5.65], size=10)
para('Дата виконання: 26 жовтня 2026 року. Повне формулювання завдання збережено в docs/technical-assignment.md.')

# Page 2: provenance, units and cleaning decisions.
heading('Дані та їх підготовка', True)
para('Джерелом є документований набір datasets::airquality зі стандартної поставки R [1]. Дані містять 153 щоденні записи за період 01.05.1973–30.09.1973. Скрипт scripts/export-data.R експортує їх у data/raw/airquality.csv; модуль аналізу читає саме цей файл функцією readr::read_csv().')
table(['Показник у модулі', 'Зміст', 'Одиниця'], [
    ['ozone_ppb','Концентрація озону','ppb'],
    ['solar_langley','Сонячна радіація','langley'],
    ['wind_ms','Швидкість вітру','м/с'],
    ['temp_c','Температура повітря','°C'],
    ['date','Дата спостереження','РРРР-ММ-ДД']
], [1.7,3.8,1.5])
para('Вихідні Wind і Temp задані у милях за годину та градусах Фаренгейта. Перетворення: wind_ms = Wind × 0,44704; temp_c = (Temp − 32) × 5/9. Month і Day використано для побудови дати, а не як неперервні пояснювальні змінні.')
table(['Результат перевірки', 'Кількість'], [
    ['Початкові рядки',audit['raw_rows']], ['Пропущені клітинки',audit['missing_cells']],
    ['Рядки з хоча б одним пропуском',audit['rows_with_missing_values']],
    ['Точні дублікати',audit['exact_duplicate_rows']],
    ['Рядки з неприпустимими значеннями',audit['invalid_rows_after_deduplication']],
    ['Рядки після очищення',audit['clean_rows']]
], [5.4,1.6])
para('Виявлено 37 пропусків Ozone та 7 пропусків Solar.R; у двох рядках відсутні обидва показники. Вилучено 42 рядки, тобто 27,45% початкової вибірки. Заповнення пропусків не проводилося. Збережені 111 рядків перевищують вимогу ТЗ щодо щонайменше 50 записів.')
para('Перевіряються схема CSV, помилки числового розбору, нескінченні й фізично неприпустимі значення, коректність календарних дат, точні дублікати та суперечливі записи за одну дату. Високі, але можливі значення не вилучаються автоматично. Повні випадки можуть відрізнятися від вилучених, тому висновки стосуються очищеної підвибірки.')

# Page 3: every requested descriptive statistic, with distribution evidence.
heading('Описова статистика та розподіли', True)
table(['Змінна', 'Середнє', 'Медіана', 'Мода', 'Частота'], [
    [names[r['variable']],num(r['mean']),num(r['median']),
     '; '.join(num(v) for v in r['modes'].split(';')),r['mode_frequency']] for r in desc
], [1.7,1.1,1.1,2.1,1.0], size=9.5)
table(['Змінна', 'Дисперсія', 'Ст відхилення', 'Мінімум', 'Максимум'], [
    [names[r['variable']],num(r['variance']),num(r['sd']),num(r['min']),num(r['max'])] for r in desc
], [1.7,1.3,1.4,1.3,1.3], size=9.5)
para('Усі показники обчислено для n = 111. Дисперсія є вибірковою, із дільником n − 1; її одиниця — квадрат одиниці відповідної змінної. Для вітру є дві рівночастотні моди. Моди описують повтори виміряних значень і залежать від точності запису.')
figure('histograms','Гістограми чотирьох числових змінних після очищення',5.5)
para('Середнє озону 42,10 ppb перевищує медіану 31 ppb. Це узгоджується з правосторонньою асиметрією та рідкісними високими концентраціями до 168 ppb. Температура має менший відносний розкид, ніж озон.')

# Page 4: both required plot families, with exploratory fit clearly named.
heading('Діаграми розмаху та розсіювання', True)
figure('boxplots','Boxplot із позначенням значень поза межами 1,5 міжквартильного розмаху',5.9)
para('Правило 1,5 IQR позначило два значення озону та три швидкості вітру. Це ознаки для перевірки, а не доказ помилки, тому спостереження збережено.')
figure('scatterplots','Озон і три метеорологічні показники; лінії є окремими однофакторними трендами',6.8)
para('Озон загалом зростає з температурою і знижується зі швидкістю вітру. Для сонячної радіації зв’язок слабший і точки помітно розсіяні. Ці графіки описують парні зв’язки, а не ізольовані причинні ефекти.')

# Page 5: correlations and significance.
heading('Кореляційний аналіз', True)
para('Для всіх 111 повних спостережень обчислено кореляції Пірсона та рангові кореляції Спірмена. Перевірено нульову гіпотезу про нульову кореляцію Пірсона. Для шести пар застосовано корекцію Benjamini–Hochberg; подані p-значення не враховують можливу часову залежність і тому мають розвідувальне значення.')
table(['Пара з озоном', 'r Пірсона', '95% ДІ', 'ρ Спірмена', 'p після BH'], [
    [names[r['variable_2']],num(r['pearson_r'],3),f"[{num(r['ci_lower'],3)}; {num(r['ci_upper'],3)}]",num(r['spearman_rho'],3),pv(r['p_adjusted_bh'])]
    for r in corr if r['variable_1']=='ozone_ppb'
], [1.7,1.0,1.8,1.2,1.3],size=9.5)
figure('correlation_heatmap','Матриця парних коефіцієнтів кореляції Пірсона',5.7)
para('Найсильніший додатний зв’язок озону спостерігається з температурою (r = 0,699); із вітром він від’ємний (r = −0,612), а з радіацією помірний (r = 0,348). Усі три зв’язки з озоном значущі на рівні 0,05 після BH. Кореляція не доводить причинності.')

# Page 6: training, coefficients, equation and fit quality.
heading('Побудова регресійної моделі', True)
para('Специфікацію ozone_ppb ~ temp_c + wind_ms + solar_langley задано до оцінювання. За датою перші 88 повних спостережень (01.05–06.09.1973) утворюють навчальну вибірку, останні 23 (07.09–30.09.1973) — тестову. Модель lm() навчається лише на перших 88 рядках; тестові значення не використовуються для оцінювання коефіцієнтів.')
para('Оцінена залежність у вихідних одиницях показників:')
math_p=doc.add_paragraph()
math_p.alignment=WD_ALIGN_PARAGRAPH.CENTER
math=OxmlElement('m:oMath'); mr=OxmlElement('m:r'); mt=OxmlElement('m:t')
mt.text='Ŷ = −0,6877 + 2,7696T − 8,3953W + 0,06572S'
mr.append(mt); math.append(mr); math_p._p.append(math)
para('Тут Ŷ — прогноз озону в ppb; T — температура в °C; W — швидкість вітру в м/с; S — сонячна радіація в langley.')
table(['Коефіцієнт', 'Оцінка', '95% ДІ', 'p'], [
    [names[r['term']],num(r['estimate'],4),f"[{num(r['ci_lower'],3)}; {num(r['ci_upper'],3)}]",pv(r['p_value'])] for r in coeff
], [1.9,1.3,2.5,1.3])
para('За фіксованих інших предикторів збільшення температури на 1 °C асоціюється зі збільшенням прогнозу озону на 2,77 ppb; збільшення вітру на 1 м/с — зі зниженням на 8,40 ppb. Зростання радіації на 100 langley відповідає збільшенню прогнозу приблизно на 6,57 ppb. Ці коефіцієнти описують умовні статистичні зв’язки.')
table(['Показник навчальної моделі', 'Значення'], [
    ['R²',num(metrics['train_r_squared'],4)], ['Скоригований R²',num(metrics['train_adjusted_r_squared'],4)],
    ['F статистика за 3 і 84 ступенів свободи',num(metrics['model_f_statistic'],3)],
    ['p загального F тесту',pv(metrics['model_f_p_value'])],
    ['RMSE на навчальній вибірці, ppb',num(metrics['train_rmse'])],
    ['Максимальний VIF','1,439']
], [5.1,1.9])
para('Усі три нахили мають p < 0,05; вільний член не є значущим. VIF 1,09–1,44 не вказують на виражену мультиколінеарність. Значущість коефіцієнтів і навчальний R² самі по собі не гарантують точності прогнозу; припущення моделі перевірено окремо.')

# Page 7: honest evaluation of held-out prediction.
heading('Перевірка прогнозування на нових датах', True)
table(['Метрика на 23 тестових днях', 'Значення'], [
    ['RMSE, ppb',num(metrics['test_rmse'])], ['MAE, ppb',num(metrics['test_mae'])],
    ['R²',num(metrics['test_r_squared'],4)],
    ['RMSE прогнозу середнім навчальної вибірки, ppb',num(metrics['test_baseline_rmse'])],
    ['RMSE відносно середнього тесту, ppb',num(metrics['test_mean_reference_rmse'])]
], [5.2,1.8])
para('Тестовий R² обчислено як 1 − SSE/SST, де SST визначено відносно середнього тестових значень. Від’ємне значення означає, що модель програє цьому ретроспективному орієнтиру. Середнє тесту заздалегідь невідоме й не є доступним практичним прогнозом. Водночас модель краща за доступний базовий прогноз середнім навчання: RMSE 18,77 проти 27,74 ppb.')
figure('test_observed_predicted','Фактичні та прогнозовані значення на відкладеній вибірці; штрихова лінія означає точну відповідність',5.9)
para('Середній озон знизився з 47,48 ppb у навчальній частині до 21,52 ppb у тестовій, а стандартне відхилення — з 35,12 до 10,01 ppb. Це свідчить про відмінність періодів і допомагає пояснити низький тестовий R². Одного пізнього тестового відрізка недостатньо для доведення сезонної причинності або стабільної якості моделі.')

# Page 8: diagnostics, plots at a legible size.
heading('Графічна перевірка залишків', True)
figure('residuals_vs_fitted','Залишки відносно прогнозів навчальної моделі',5.8)
figure('residual_qq','Нормальний Q–Q графік залишків',5.8)
para('На графіках перевіряються систематична структура, неоднаковий розкид і відхилення від нормального розподілу. Відхилення верхнього хвоста Q–Q графіка узгоджується з результатом тесту Шапіро–Вілка.')

# Page 9: diagnostic numbers and limitations.
heading('Діагностика та межі застосування', True)
table(['Перевірка', 'Результат', 'Інтерпретація'], [
    ['Шапіро–Вілк',f"W = {num(diagnostics['shapiro_w'],3)}; p = {pv(diagnostics['shapiro_p_value'])}",'Нормальність залишків відхиляється на рівні 0,05.'],
    ['Допоміжний LM тест гетероскедастичності',f"LM = {num(diagnostics['heteroskedasticity_lm_statistic'],3)}; p = {num(diagnostics['heteroskedasticity_p_value'],3)}",'У цьому тесті немає підстав відхилити сталу дисперсію; це не доведення її сталості.'],
    ['Послідовність залишків',f"r₁ = {num(diagnostics['residual_lag1_correlation'],3)}; DW = {num(diagnostics['durbin_watson_statistic'],3)}",'Описові індикатори близькі до відсутності залежності сусідніх залишків; формального p немає.'],
    ['Вплив спостережень',f"Cook max = {num(diagnostics['max_cooks_distance'],3)}; позначено 5",'Перевищено орієнтир 4/88 = 0,0455; рядки збережено.']
], [2.0,2.0,3.0],size=10)
para('LM діагностика використовує допоміжну регресію квадратів залишків на прогнозовані значення, статистику nR² та один ступінь свободи. Вона не перевіряє всі можливі форми неоднорідності. Через пропуски календарних днів лаг 1 означає сусідні доступні спостереження, а не завжди сусідні дати.')
para('Ненормальність залишків і впливові спостереження зменшують надійність точних t- та F-висновків і нормальних інтервалів прогнозу. Низький тестовий R² підтверджує, що модель слід використовувати як навчальний опис залежностей, а не як готову систему прогнозування якості повітря.')
doc.add_heading('Що потрібно для практичного застосування',2)
para('Доцільно перевірити перетворення цільової змінної, нелінійні залежності та взаємодії, додати часові або сезонні предиктори й порівняти моделі за послідовною часовою валідацією. Такий добір треба проводити на навчальних періодах, зберігаючи окремі дані для остаточної перевірки.')
para('Дані історичні й локальні. Механізм пропусків не встановлено, 27,45% записів вилучено, а тест містить лише 23 дні. Перенесення висновків на сучасні міста потребує нових даних, перевірки одиниць вимірювання та повторної валідації.')

# Page 10: genuine scenario forecasts with intervals.
heading('Прогнозні сценарії та інтервали', True)
para('Для демонстрації predict() задано три сценарії в межах діапазонів навчальних предикторів. Це перевірка кожної змінної окремо, а не гарантія достатнього покриття спільного простору ознак. Інтервали мають номінальний рівень 95% і стосуються окремого майбутнього спостереження.')
table(['Сценарій','T °C','W м/с','S langley','Прогноз ppb','95% інтервал ppb'], [
    [r['scenario'],num(r['temp_c'],0),num(r['wind_ms'],0),num(r['solar_langley'],0),num(r['fit']),f"[{num(r['lwr'])}; {num(r['upr'])}]"] for r in scenarios
], [0.8,.7,.8,1.0,1.1,2.6],size=9.5)
figure('test_forecast','Хронологічні прогнози та 95% прогнозні інтервали на тестових днях',6.4)
para('Усі 23 тестові значення потрапили в номінальні 95% інтервали. За такої малої вибірки та широких інтервалів покриття 100% не доводить їхньої каліброваності. Для сценарію A нижня межа дорівнює −13,58 ppb, що фізично неможливо. Межу показано без прихованого обрізання: це прояв обмеження нормальної лінійної моделі.')
para('Порівняння B та C демонструє умовний ефект швидкості вітру: її зменшення з 4 до 3 м/с за однакових температури й радіації збільшує модельний прогноз приблизно на 8,40 ppb. Інтервали залишаються широкими, тому такі числа не слід трактувати як точні практичні передбачення.')

# Pages 11 and 12: genuine application screenshots, not simulated terminals.
heading('Робочий код у RStudio', True)
para('Проєкт містить analysis.R із функцією run_analysis(), вихідний CSV, файл Practice_R.Rproj та окремі каталоги результатів. Основні операції реалізовано через readr, dplyr, ggplot2 і вбудовані функції R. На знімку показано фактичний файл коду у запущеному RStudio Server.')
picture(OUT/'screenshots/rstudio-code.png','Код аналізу у редакторі RStudio',6.85)
doc.add_heading('Відтворення повного запуску',2)
code('bash scripts/setup-r.sh\nsource scripts/activate-r.sh\nRscript analysis.R\nRscript tests/test-analysis.R')
para('У RStudio відкрийте Practice_R.Rproj та виконайте source("analysis.R"). Для іншого CSV з тією самою схемою використовуйте Rscript analysis.R шлях/до/input.csv каталог_результатів. Схема: Ozone, Solar.R, Wind, Temp, Month, Day; одиниці відповідають вихідному airquality.')

heading('Фактичне виконання операцій', True)
para('Наведений знімок отримано після виконання source("analysis.R") у RStudio. Числові результати з консолі збігаються з CSV-таблицями, використаними у звіті. Знімки показують робочий інтерфейс і фактичні результати запуску.')
picture(OUT/'screenshots/rstudio-results.png','Результати статистичного аналізу в консолі RStudio',6.85)
table(['Перевірка модуля','Підсумок'],[
    ['Невідома схема CSV','Вхідні дані відхилено з поясненням.'],
    ['Некоректне числове значення','Помилку розбору зареєстровано; аналіз зупинено.'],
    ['Менше 50 повних записів','Вхідний набір відхилено.'],
    ['Дублікати, неможливі дати, від’ємний озон','Контрольні помилки виявлено; коректні рядки збережено.'],
    ['Відновлення збереженої моделі','Прогнози з RDS збігаються з таблицею; межі інтервалів коректні.']
], [3.1,3.9],size=10)
para('Усі п’ять інтеграційних перевірок пройдено. Вони виконуються на тимчасових даних і не підміняють результати основного дослідження.')

# Page 13: conclusions, handover and sources.
heading('Висновки та склад результатів', True)
para('Виконано всі сім етапів технічного завдання: підготовку даних, описову статистику, візуалізацію, кореляційний і регресійний аналіз, перевірку моделі та інтерпретацію. Практичний результат — програмний модуль, очищений CSV, статистичні таблиці, вісім графіків, збережена модель і цей звіт у Word та PDF.')
para('У дослідженій вибірці підвищена концентрація озону пов’язана з вищою температурою та нижчою швидкістю вітру. Після спільного врахування трьох предикторів усі нахили статистично значущі, а навчальний R² становить 0,6197. Цей результат характеризує наявні історичні дані й не є причинним висновком.')
para('Прогностична перевірка виявила суттєві обмеження: тестовий R² від’ємний, залишки ненормальні, є впливові спостереження, а інтервали широкі. Модель корисна для демонстрації статистичного процесу та формування гіпотез. Для практичного прогнозування потрібні додаткові дані та валідація вдосконалених специфікацій.')
doc.add_heading('Основні файли',2)
table(['Артефакт','Розташування'],[
    ['Код та проєкт','analysis.R; Practice_R.Rproj'],
    ['Вихідні та очищені дані','data/raw/airquality.csv; data/processed/airquality_clean.csv'],
    ['Таблиці та графіки','results/tables/; results/figures/'],
    ['Модель','results/models/ozone_linear_model.rds'],
    ['Журнал і версії','results/analysis-output.txt; results/session-info.txt'],
    ['Перевірки','tests/test-analysis.R'],
    ['Звіт та знімки','reports/Practice_R_Report.docx; reports/Practice_R_Report.pdf; reports/screenshots/']
], [2.1,4.9],size=9.5)
doc.add_heading('Використані джерела',2)
for text in [
    '[1] R Documentation. airquality — New York Air Quality Measurements. https://stat.ethz.ch/R-manual/R-devel/library/datasets/html/airquality.html',
    '[2] R Documentation. lm, summary.lm, predict.lm, cor.test, shapiro.test. https://stat.ethz.ch/R-manual/R-devel/library/stats/html/00Index.html',
    '[3] Документація пакетів ggplot2, dplyr та readr. https://ggplot2.tidyverse.org/ ; https://dplyr.tidyverse.org/ ; https://readr.tidyverse.org/',
    '[4] Репозиторій програми та результатів. https://github.com/VladyslavGennadijovych/Practice_R'
]:
    p=para(text)
    for r in p.runs: r.font.size=Pt(9)

# The bundled Word template may carry a decorative Title paragraph border.
for root in [doc.styles.element, doc.element]:
    for border in list(root.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
doc.save(OUT/'Practice_R_Report.docx')
print(OUT/'Practice_R_Report.docx')
